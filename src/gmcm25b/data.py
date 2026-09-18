"""Streaming, traceable CSI ingestion for question 1.

The paper's appendix divides channel power by 1000; its main-text expression
does not. ``signal_scale`` exposes that distinction. Noise epsilon is in watts
per subcarrier and defaults to zero because the paper gives no numeric value.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
from typing import Any, Sequence
from zipfile import BadZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook


N_SUBCARRIERS = 122
CSI_COLUMNS = tuple(f"csi_matrix_r{r}_c{c}" for r in range(2) for c in range(4))
META_COLUMNS = (
    "sample_id", "source_file", "source_index", "excel_row", "device", "split",
    "status", "reason", "noise_floor", "beamforming_en", "dfx_time", "csi_time",
)
CACHE_SCHEMA = 1


@dataclass
class Dataset:
    """Rows of linear SINR and aligned, explicit source metadata."""

    sinr: np.ndarray
    meta: pd.DataFrame
    y: np.ndarray | None = None

    def __post_init__(self) -> None:
        if self.sinr.ndim != 2 or self.sinr.shape[1] != N_SUBCARRIERS:
            raise ValueError(f"SINR must have shape (n, {N_SUBCARRIERS})")
        if len(self.meta) != len(self.sinr):
            raise ValueError("Metadata and SINR row counts differ")
        if self.y is not None and self.y.shape != (len(self.sinr),):
            raise ValueError("Labels must have shape (n,)")

    def __len__(self) -> int:
        return len(self.sinr)


def parse_csi(value: Any, expected_size: int = N_SUBCARRIERS) -> np.ndarray:
    """Parse a list of finite complex literals, without evaluating Python code."""
    if not isinstance(value, str):
        raise ValueError("CSI must be a string containing a complex-number list")
    text = value.strip()
    if not (text.startswith("[") and text.endswith("]")):
        raise ValueError("CSI must be enclosed in square brackets")
    tokens = text[1:-1].split(",")
    if len(tokens) != expected_size:
        raise ValueError(f"Expected {expected_size} CSI values; got {len(tokens)}")
    try:
        result = np.fromiter((complex(token.strip()) for token in tokens),
                             dtype=np.complex128, count=expected_size)
    except (ValueError, TypeError, OverflowError) as error:
        raise ValueError("Invalid complex literal in CSI") from error
    if not np.isfinite(result).all():
        raise ValueError("CSI contains a non-finite value")
    return result


def _check_settings(signal_scale: float, noise_epsilon_w: float) -> None:
    if not np.isfinite(signal_scale) or signal_scale <= 0:
        raise ValueError("signal_scale must be finite and positive")
    if not np.isfinite(noise_epsilon_w) or noise_epsilon_w < 0:
        raise ValueError("noise_epsilon_w must be finite and nonnegative")


def channel_to_sinr(
    channel: np.ndarray,
    noise_floor_dbm: float,
    signal_scale: float = 0.001,
    noise_epsilon_w: float = 0.0,
) -> np.ndarray:
    """Convert an (8,122) or (2,4,122) channel to 122 linear SINR values."""
    _check_settings(signal_scale, noise_epsilon_w)
    values = np.asarray(channel, dtype=np.complex128)
    if values.shape not in ((8, N_SUBCARRIERS), (2, 4, N_SUBCARRIERS)):
        raise ValueError("Channel must have shape (8,122) or (2,4,122)")
    if not np.isfinite(values).all():
        raise ValueError("Channel contains non-finite values")
    noise = float(noise_floor_dbm)
    if not np.isfinite(noise) or not -120 <= noise <= -50:
        raise ValueError("Noise floor must be between -120 and -50 dBm")
    noise_per_subcarrier_w = 10.0 ** ((noise - 30.0) / 10.0) / N_SUBCARRIERS
    with np.errstate(over="ignore", invalid="ignore"):
        power = np.sum(np.abs(values.reshape(8, N_SUBCARRIERS)) ** 2, axis=0)
        sinr = signal_scale * power / (noise_per_subcarrier_w + noise_epsilon_w)
    if not np.isfinite(sinr).all():
        raise ValueError("Computed SINR is non-finite")
    return sinr


def _number(value: Any) -> float:
    if isinstance(value, (bool, np.bool_)) or value is None:
        return float("nan")
    try:
        return float(value)
    except (TypeError, ValueError, OverflowError):
        return float("nan")


def _summary(dataset: Dataset) -> dict[str, Any]:
    reasons: Counter[str] = Counter()
    for reason in dataset.meta["reason"]:
        reasons.update(str(reason).split(";") if reason else ())
    ok = int(dataset.meta["status"].eq("ok").sum())
    return {"rows_total": len(dataset), "rows_ok": ok,
            "rows_invalid": len(dataset) - ok, "reason_counts": dict(reasons)}


def read_workbook(
    path: str | Path,
    *,
    split: str,
    source_file: str | None = None,
    device: str | None = None,
    signal_scale: float = 0.001,
    noise_epsilon_w: float = 0.0,
) -> Dataset:
    """Read every workbook row; invalid rows stay aligned with all-NaN SINR.

    ``prepare_dataset`` filters invalid training rows afterwards. Validation
    rows are never dropped. Worksheet row numbers include the header at row 1.
    """
    _check_settings(signal_scale, noise_epsilon_w)
    if split not in ("train", "valid"):
        raise ValueError("split must be train or valid")
    path = Path(path)
    source_file = source_file or path.as_posix()
    device = device or path.parent.parent.name
    workbook = load_workbook(path, read_only=True, data_only=True)
    sinr_rows: list[np.ndarray] = []
    rows: list[dict[str, Any]] = []
    labels: list[float] = []
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            raise ValueError(f"Workbook is empty: {path}")
        named = [str(name) for name in header if name is not None]
        if len(named) != len(set(named)):
            raise ValueError(f"Duplicate named column in {path}")
        positions = {str(name): index for index, name in enumerate(header)
                     if name is not None}
        required = (*CSI_COLUMNS, "beamforming_en", "noise_floor")
        if split == "train":
            required = (*required, "mcs")
        missing = [name for name in required if name not in positions]
        if missing:
            raise ValueError(f"Missing columns in {path}: {', '.join(missing)}")

        for excel_row, row in enumerate(iterator, start=2):
            def cell(name: str) -> Any:
                index = positions.get(name)
                return row[index] if index is not None and index < len(row) else None

            reasons: list[str] = []
            noise = _number(cell("noise_floor"))
            beamforming = _number(cell("beamforming_en"))
            if not np.isfinite(noise) or not -120 <= noise <= -50:
                reasons.append("invalid_noise_floor")
            if not np.isfinite(beamforming) or beamforming != 0:
                reasons.append("beamforming_not_zero")
            label = _number(cell("mcs")) if split == "train" else float("nan")
            if split == "train" and (not np.isfinite(label) or label < 0):
                reasons.append("invalid_mcs")
            channels = []
            for column in CSI_COLUMNS:
                try:
                    channels.append(parse_csi(cell(column)))
                except ValueError:
                    reasons.append(f"invalid_csi:{column}")
            sinr = np.full(N_SUBCARRIERS, np.nan)
            if not reasons:
                try:
                    sinr = channel_to_sinr(np.stack(channels), noise,
                                           signal_scale, noise_epsilon_w)
                except ValueError:
                    reasons.append("invalid_computed_sinr")
            source_index = str(row[0]) if row and row[0] is not None else ""
            rows.append({
                "sample_id": f"{source_file}::{excel_row}",
                "source_file": source_file, "source_index": source_index,
                "excel_row": excel_row, "device": device, "split": split,
                "status": "invalid" if reasons else "ok",
                "reason": ";".join(reasons), "noise_floor": cell("noise_floor"),
                "beamforming_en": cell("beamforming_en"),
                "dfx_time": cell("dfx_time"), "csi_time": cell("csi_time"),
            })
            sinr_rows.append(sinr)
            labels.append(label)
    finally:
        workbook.close()
    matrix = np.stack(sinr_rows) if sinr_rows else np.empty((0, N_SUBCARRIERS))
    return Dataset(matrix, pd.DataFrame(rows, columns=META_COLUMNS),
                   np.asarray(labels) if split == "train" else None)


def prepare_dataset(datasets: Sequence[Dataset]) -> tuple[Dataset, Dataset, dict[str, Any]]:
    """Combine sources, filter training failures, and preserve every valid row."""
    groups: dict[str, list[Dataset]] = {"train": [], "valid": []}
    for dataset in datasets:
        if len(dataset):
            split_values = dataset.meta["split"].unique()
            if len(split_values) != 1 or split_values[0] not in groups:
                raise ValueError("Each dataset must contain exactly one known split")
            groups[str(split_values[0])].append(dataset)
    result: dict[str, Dataset] = {}
    audit: dict[str, Any] = {}
    for split, pieces in groups.items():
        matrix = np.concatenate([piece.sinr for piece in pieces]) if pieces else np.empty((0, N_SUBCARRIERS))
        meta = pd.concat([piece.meta for piece in pieces], ignore_index=True) if pieces else pd.DataFrame(columns=META_COLUMNS)
        if split == "train" and any(piece.y is None for piece in pieces):
            raise ValueError("Training labels are missing")
        y = np.concatenate([piece.y for piece in pieces]) if split == "train" and pieces else (np.empty(0) if split == "train" else None)
        combined = Dataset(matrix, meta, y)
        audit[split] = _summary(combined)
        if split == "train":
            keep = meta["status"].eq("ok").to_numpy(dtype=bool)
            result[split] = Dataset(matrix[keep], meta.loc[keep].reset_index(drop=True), y[keep])
            audit[split]["rows_retained"] = int(keep.sum())
            audit[split]["rows_dropped"] = int((~keep).sum())
        else:
            result[split] = combined
            audit[split]["rows_retained"] = len(combined)
            audit[split]["rows_dropped"] = 0
    return result["train"], result["valid"], audit


def file_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_json(path: Path, data: dict[str, Any]) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_dataset(
    raw_root: str | Path,
    manifest_path: str | Path,
    cache_dir: str | Path,
    signal_scale: float = 0.001,
    noise_epsilon_w: float = 0.0,
) -> tuple[Dataset, Dataset, dict[str, Any]]:
    """Load manifest sources with SHA256- and setting-bound per-file caches.

    The row audit CSV includes discarded training rows and all validation rows.
    Cache metadata is written last; interrupted writes are recomputed on retry.
    No pickle loading or raw spreadsheet modification is used.
    """
    _check_settings(signal_scale, noise_epsilon_w)
    raw_root = Path(raw_root).resolve()
    cache_dir = Path(cache_dir)
    cache_dir.mkdir(parents=True, exist_ok=True)
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8-sig"))
    settings = {"cache_schema": CACHE_SCHEMA, "signal_scale": signal_scale,
                "noise_epsilon_w": noise_epsilon_w, "n_subcarriers": N_SUBCARRIERS,
                "noise_dbm_range": [-120, -50], "beamforming_en": 0,
                "mcs_requirement": "finite and nonnegative"}
    datasets: list[Dataset] = []
    source_audits: list[dict[str, Any]] = []
    seen: set[str] = set()
    for entry in manifest["files"]:
        relative = str(entry["path"]).replace("\\", "/")
        if relative in seen:
            raise ValueError(f"Duplicate source in manifest: {relative}")
        seen.add(relative)
        path = (raw_root / relative).resolve()
        if not path.is_relative_to(raw_root):
            raise ValueError(f"Manifest path escapes raw_root: {relative}")
        if not path.is_file():
            raise FileNotFoundError(f"Dataset file is missing: {path}")
        size = path.stat().st_size
        if entry.get("size_bytes") is not None and size != int(entry["size_bytes"]):
            raise ValueError(f"Incomplete or changed file {relative}: expected {entry['size_bytes']} bytes, found {size}")
        sha256 = file_sha256(path)
        if entry.get("sha256") and sha256 != entry["sha256"]:
            raise ValueError(f"SHA256 mismatch: {relative}")
        identity = {"path": relative, "sha256": sha256, "device": entry["device"],
                    "split": entry["split"], "settings": settings}
        key = hashlib.sha256(json.dumps(identity, sort_keys=True).encode()).hexdigest()
        numeric_path = cache_dir / f"{key}.npz"
        meta_path = cache_dir / f"{key}.csv"
        info_path = cache_dir / f"{key}.json"
        dataset = None
        if numeric_path.exists() and meta_path.exists() and info_path.exists():
            try:
                info = json.loads(info_path.read_text(encoding="utf-8"))
                if info["identity"] == identity:
                    meta = pd.read_csv(meta_path, keep_default_na=False,
                                       dtype={"source_index": str, "sample_id": str,
                                              "source_file": str, "device": str})
                    with np.load(numeric_path, allow_pickle=False) as arrays:
                        dataset = Dataset(arrays["sinr"], meta,
                                          arrays["y"] if entry["split"] == "train" else None)
            except (OSError, ValueError, KeyError, EOFError, BadZipFile):
                dataset = None
        if dataset is None:
            print(f"Reading {relative}", flush=True)
            dataset = read_workbook(path, split=entry["split"], source_file=relative,
                                    device=entry["device"], signal_scale=signal_scale,
                                    noise_epsilon_w=noise_epsilon_w)
            dataset.meta.to_csv(meta_path, index=False, encoding="utf-8-sig")
            np.savez_compressed(numeric_path, sinr=dataset.sinr,
                                y=dataset.y if dataset.y is not None else np.empty(0))
            _write_json(info_path, {"identity": identity, "summary": _summary(dataset)})
        datasets.append(dataset)
        source_audits.append({"path": relative, "device": entry["device"],
                              "split": entry["split"], "size_bytes": size,
                              "sha256": sha256, **_summary(dataset)})
    train, valid, audit = prepare_dataset(datasets)
    row_audit = pd.concat([dataset.meta for dataset in datasets], ignore_index=True) if datasets else pd.DataFrame(columns=META_COLUMNS)
    all_key = hashlib.sha256(json.dumps({"sources": source_audits, "settings": settings}, sort_keys=True).encode()).hexdigest()[:16]
    row_audit_path = cache_dir / f"row_audit_{all_key}.csv"
    row_audit.to_csv(row_audit_path, index=False, encoding="utf-8-sig")
    audit.update({"settings": settings, "sources": source_audits,
                  "row_audit_csv": str(row_audit_path.resolve()),
                  "valid_order": "manifest file order, then original worksheet row order"})
    _write_json(cache_dir / f"audit_{all_key}.json", audit)
    return train, valid, audit

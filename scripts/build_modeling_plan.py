"""Reconstruct the reviewed modeling proposal and build an editable Word document.

This script is executed by the repository's GitHub Actions workflow. The temporary
base64 source fragments in .build are checked against the original source hash.
"""
from pathlib import Path
import base64
import bz2
import hashlib
import subprocess
import zipfile

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
PARTS = [ROOT / '.build' / f'source.part{i:02d}' for i in range(1, 5)]
OUT = ROOT / 'docs' / '2025B-建模思路与实施方案.docx'
B64_SHA256 = '30cd5b22fec06260bdf1849512f3e528ac19ab2a572962e9fd3799d49b3d57c5'
MD_SHA256 = '19f6fcd0dc26c0f215ea3551eb67c86996c14485ed1c1a3799fb92e9d7806ded'


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def main():
    payload = ''.join(path.read_text(encoding='ascii').strip() for path in PARTS).encode('ascii')
    if sha256(payload) != B64_SHA256:
        raise RuntimeError('Source fragment checksum mismatch: do not commit an incomplete proposal')
    md = bz2.decompress(base64.b64decode(payload, validate=True))
    if sha256(md) != MD_SHA256:
        raise RuntimeError('Decompressed Markdown checksum mismatch')

    OUT.parent.mkdir(parents=True, exist_ok=True)
    source = ROOT / '.build' / 'source-reconstructed.md'
    source.write_bytes(md)
    subprocess.run(['pandoc', str(source), '-f', 'markdown+tex_math_dollars', '-t', 'docx', '-o', str(OUT)], check=True)

    doc = Document(OUT)
    section = doc.sections[0]
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.1)
    section.left_margin, section.right_margin = Cm(2.6), Cm(2.4)
    styles = doc.styles
    for name in ('Normal', 'Body Text', 'First Paragraph'):
        try:
            style = styles[name]
        except KeyError:
            continue
        style.font.name = '宋体'
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor(23, 29, 39)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
        style.paragraph_format.line_spacing = 1.35
        style.paragraph_format.space_after = Pt(5)
    for name, size in (('Title', 18), ('Heading1', 14), ('Heading2', 12), ('Heading3', 11)):
        try:
            style = styles[name]
        except KeyError:
            continue
        style.font.name = '黑体'
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(20, 33, 51)
        style.paragraph_format.space_before = Pt(0 if name == 'Title' else 14)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    for name in ('Author', 'Date'):
        try:
            style = styles[name]
        except KeyError:
            continue
        style.font.size = Pt(10)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    for index, para in enumerate(doc.paragraphs):
        if index == 0 or para.style.name in ('Author', 'Date'):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if para.style.name in ('Body Text', 'First Paragraph'):
            para.paragraph_format.first_line_indent = Cm(0.75)
        if para.style.name.startswith('Heading'):
            para.paragraph_format.keep_together = True
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('— ')
    run = footer.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    run._r.append(fld_begin)
    run = footer.add_run()
    fld_text = OxmlElement('w:instrText')
    fld_text.set(qn('xml:space'), 'preserve')
    fld_text.text = ' PAGE '
    run._r.append(fld_text)
    run = footer.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_end)
    footer.add_run(' —')
    doc.core_properties.title = '2025 B题无线通信系统链路速率建模：求解思路与实施方案'
    doc.core_properties.subject = '建模思路、三问方法、实验验证及实施路线'
    doc.core_properties.author = 'GMCM-25B-Reproduction'
    doc.save(OUT)
    with zipfile.ZipFile(OUT) as package:
        assert package.testzip() is None
        assert b'<m:oMath' in package.read('word/document.xml'), 'Word equations missing'
    print(f'Word created and verified: {OUT}, {OUT.stat().st_size} bytes')


if __name__ == '__main__':
    main()
